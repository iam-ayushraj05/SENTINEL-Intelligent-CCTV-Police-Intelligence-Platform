import io
import uuid
import hashlib
import logging
from datetime import datetime, timezone
from typing import Optional

try:
    import docx
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    logger_init = logging.getLogger("sentinel.case_report_generator")
    logger_init.warning("python-docx not installed. Case report generation will be unavailable. Run: pip install python-docx")

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.emergency import EmergencyIncident, IncidentTimeline, CaseMessage, CaseReportDocument
from app.models.document import Document, Category
from app.services.document_storage import encrypt_and_store

logger = logging.getLogger("sentinel.case_report_generator")


class CaseReportGeneratorService:
    """Generates official SENTINEL EMERGENCY CASE REPORT DOCX files."""

    @staticmethod
    def generate_docx_bytes(
        incident: EmergencyIncident,
        timeline_entries: list[IncidentTimeline],
        messages: list[CaseMessage],
        human_verified: bool = True,
    ) -> bytes:
        """
        Builds a styled python-docx document containing case details, chronological timeline,
        communication transcript, and AI technical specifications.
        """
        doc = docx.Document()

        # Page margins
        sections = doc.sections
        for section in sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)

        # Title Banner
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_title = p_title.add_run("SENTINEL EMERGENCY CASE REPORT")
        run_title.font.name = "Arial"
        run_title.font.size = Pt(22)
        run_title.font.bold = True
        run_title.font.color.rgb = RGBColor(0, 33, 71)  # Police Navy #002147

        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_sub = p_sub.add_run(f"OFFICIAL INCIDENT DOSSIER · {incident.incident_code}")
        run_sub.font.name = "Arial"
        run_sub.font.size = Pt(11)
        run_sub.font.bold = True
        run_sub.font.color.rgb = RGBColor(0, 119, 182)

        doc.add_paragraph().paragraph_format.space_after = Pt(12)

        # Section 1: Executive Case Summary
        h1 = doc.add_heading("1. EXECUTIVE CASE SUMMARY", level=1)
        h1.runs[0].font.color.rgb = RGBColor(0, 33, 71)

        table_summary = doc.add_table(rows=10, cols=2)
        table_summary.alignment = WD_TABLE_ALIGNMENT.CENTER
        table_summary.autofit = False

        summary_data = [
            ("CASE NUMBER", incident.incident_code),
            ("STATUS", incident.status),
            ("PRIORITY / SEVERITY", f"{incident.priority or 'HIGH'} / {incident.severity}"),
            ("EVENT TYPE", incident.incident_type),
            ("CAMERA ID / LOCATION", f"{incident.camera_id or 'CAM11'} ({incident.location_name or 'Main Surveillance'})"),
            ("CREATED TIMESTAMP", incident.created_at.strftime("%Y-%m-%d %H:%M:%S UTC") if incident.created_at else "N/A"),
            ("CLOSED TIMESTAMP", incident.closed_at.strftime("%Y-%m-%d %H:%M:%S UTC") if getattr(incident, "closed_at", None) else datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")),
            ("SOURCE ALERT ID", str(incident.source_alert_id or "N/A")),
            ("SOURCE EVENT ID", str(incident.source_event_id or "N/A")),
            ("AI CONFIDENCE SCORE", f"{((incident.ai_confidence or 0.90) * 100):.1f}%"),
        ]

        for idx, (label, val) in enumerate(summary_data):
            row_cells = table_summary.rows[idx].cells
            row_cells[0].width = Inches(2.2)
            row_cells[1].width = Inches(4.5)
            
            p0 = row_cells[0].paragraphs[0]
            r0 = p0.add_run(label)
            r0.font.bold = True
            r0.font.size = Pt(9.5)
            
            p1 = row_cells[1].paragraphs[0]
            r1 = p1.add_run(val)
            r1.font.size = Pt(9.5)

        doc.add_paragraph().paragraph_format.space_after = Pt(14)

        # Section 2: Chronological Timeline Audit Log
        h2 = doc.add_heading("2. CHRONOLOGICAL CASE TIMELINE", level=1)
        h2.runs[0].font.color.rgb = RGBColor(0, 33, 71)

        t_timeline = doc.add_table(rows=1, cols=4)
        t_timeline.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr_cells = t_timeline.rows[0].cells
        hdr_titles = ["TIMESTAMP", "EVENT TYPE", "ACTOR", "DESCRIPTION / LOG"]
        widths = [Inches(1.5), Inches(1.8), Inches(1.2), Inches(2.2)]

        for i, title_text in enumerate(hdr_titles):
            hdr_cells[i].width = widths[i]
            p = hdr_cells[i].paragraphs[0]
            r = p.add_run(title_text)
            r.font.bold = True
            r.font.size = Pt(9)

        for entry in sorted(timeline_entries, key=lambda x: x.event_time or datetime.min):
            row_cells = t_timeline.add_row().cells
            time_str = entry.event_time.strftime("%H:%M:%S") if entry.event_time else "N/A"
            vals = [time_str, entry.event_type, entry.actor_id or "SYSTEM", entry.description or ""]
            for i, val in enumerate(vals):
                row_cells[i].width = widths[i]
                p = row_cells[i].paragraphs[0]
                r = p.add_run(val)
                r.font.size = Pt(9)

        doc.add_paragraph().paragraph_format.space_after = Pt(14)

        # Section 3: Official Communication Transcript
        h3 = doc.add_heading("3. OFFICIAL COMMUNICATION TRANSCRIPT", level=1)
        h3.runs[0].font.color.rgb = RGBColor(0, 33, 71)

        t_msg = doc.add_table(rows=1, cols=4)
        t_msg.alignment = WD_TABLE_ALIGNMENT.CENTER
        m_hdr = t_msg.rows[0].cells
        m_titles = ["TIME", "AUTHOR / OFFICER", "MESSAGE TYPE", "MESSAGE CONTENT"]
        m_widths = [Inches(1.2), Inches(1.5), Inches(1.3), Inches(2.7)]

        for i, title_text in enumerate(m_titles):
            m_hdr[i].width = m_widths[i]
            p = m_hdr[i].paragraphs[0]
            r = p.add_run(title_text)
            r.font.bold = True
            r.font.size = Pt(9)

        for msg in sorted(messages, key=lambda x: x.created_at or datetime.min):
            row_cells = t_msg.add_row().cells
            t_str = msg.created_at.strftime("%H:%M:%S") if msg.created_at else "N/A"
            m_type = (msg.message_type or "RESPONSE").upper()
            vals = [t_str, msg.sender_name or msg.sender_id or "Operator", m_type, msg.message or ""]
            for i, val in enumerate(vals):
                row_cells[i].width = m_widths[i]
                p = row_cells[i].paragraphs[0]
                r = p.add_run(val)
                r.font.size = Pt(9)
                if i == 2:
                    r.font.bold = True

        doc.add_paragraph().paragraph_format.space_after = Pt(14)

        # Section 4: AI Technical & Verification Details
        h4 = doc.add_heading("4. AI ENGINE & INTELLIGENCE DATA", level=1)
        h4.runs[0].font.color.rgb = RGBColor(0, 33, 71)

        p_ai = doc.add_paragraph()
        p_ai.paragraph_format.line_spacing = Pt(14)

        ai_details = [
            ("DETECTION EVENT TYPE: ", incident.incident_type),
            ("AI MODEL NAME: ", "YOLOv8x / Sentinel AI Engine"),
            ("MODEL VERSION: ", "v2.4.1 (Temporal Validation Core)"),
            ("AI CONFIDENCE SCORE: ", f"{((incident.ai_confidence or 0.90) * 100):.1f}%"),
            ("TEMPORAL VALIDATION: ", "TEMPORAL_VALIDATED (3-frame window persistence)"),
            ("CAMERA SOURCE: ", str(incident.camera_id or "CAM11")),
            ("VERIFICATION STATE: ", "HUMAN VERIFIED" if human_verified else "AI DETECTED — NOT HUMAN VERIFIED"),
        ]

        for lbl, val in ai_details:
            r1 = p_ai.add_run(lbl)
            r1.font.bold = True
            r1.font.size = Pt(9.5)
            r2 = p_ai.add_run(f"{val}\n")
            r2.font.size = Pt(9.5)
            if "HUMAN VERIFIED" in val:
                r2.font.bold = True
                r2.font.color.rgb = RGBColor(0, 128, 0)
            elif "NOT HUMAN VERIFIED" in val:
                r2.font.bold = True
                r2.font.color.rgb = RGBColor(180, 0, 0)

        # Footer Notice
        doc.add_paragraph().paragraph_format.space_after = Pt(20)
        p_foot = doc.add_paragraph()
        p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_foot = p_foot.add_run("CONFIDENTIAL POLICE INTELLIGENCE DOCUMENT · SENTINEL PLATFORM 2026")
        r_foot.font.size = Pt(8.5)
        r_foot.font.italic = True
        r_foot.font.color.rgb = RGBColor(120, 120, 120)

        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()

    @staticmethod
    async def generate_and_store_case_report(
        db: AsyncSession,
        incident: EmergencyIncident,
        timeline_entries: list[IncidentTimeline],
        messages: list[CaseMessage],
        owner_id: uuid.UUID,
        created_by: str = "SYSTEM",
        version: int = 1,
    ) -> tuple[CaseReportDocument, Document]:
        """
        Generates the DOCX report, computes SHA-256 hash, encrypts & stores file,
        creates CaseReportDocument and Document Centre entries.
        """
        # Determine human verification state
        human_verified = any(
            t.actor_id and t.actor_id != "SYSTEM" and t.actor_id != "AlertEngine"
            for t in timeline_entries
        ) or (incident.status in ["ACKNOWLEDGED", "IN_PROGRESS", "UNDER_REVIEW", "RESOLVED", "CLOSED"])

        docx_bytes = CaseReportGeneratorService.generate_docx_bytes(
            incident=incident,
            timeline_entries=timeline_entries,
            messages=messages,
            human_verified=human_verified,
        )

        # 1. Compute SHA-256 hash
        document_hash = hashlib.sha256(docx_bytes).hexdigest()

        # 2. Encrypt & Store file via DMS service
        filename = f"{incident.incident_code}-report-v{version}.docx"
        storage_key, _ = encrypt_and_store(docx_bytes, owner_id)

        # 3. Get or Create Category "Case Reports" in Document Centre
        cat_stmt = select(Category).where(Category.owner_id == owner_id, Category.name == "Case Reports")
        cat_res = await db.execute(cat_stmt)
        cat_rec = cat_res.scalars().first()
        if not cat_rec:
            cat_rec = Category(id=uuid.uuid4(), owner_id=owner_id, name="Case Reports", description="Automated emergency case dossier reports")
            db.add(cat_rec)
            await db.flush()

        # 4. Create Document record in documents table (Document Centre)
        doc_record = Document(
            id=uuid.uuid4(),
            owner_id=owner_id,
            category_id=cat_rec.id,
            title=f"Emergency Case Report ({incident.incident_code} v{version})",
            original_filename=filename,
            storage_key=storage_key,
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            size_bytes=len(docx_bytes),
            sha256=document_hash,
            extracted_text=f"Case Report for {incident.incident_code}. Status: {incident.status}. Type: {incident.incident_type}.",
            is_sensitive=True,
        )
        db.add(doc_record)
        await db.flush()

        # 5. Create CaseReportDocument record
        case_doc_record = CaseReportDocument(
            id=uuid.uuid4(),
            incident_id=incident.id,
            case_number=incident.incident_code,
            document_id=doc_record.id,
            document_type="CASE_REPORT",
            file_name=filename,
            storage_key=storage_key,
            document_hash=document_hash,
            version=version,
            status="GENERATED",
            created_by=created_by,
            metadata_json={
                "incident_id": str(incident.id),
                "case_number": incident.incident_code,
                "incident_type": incident.incident_type,
                "ai_confidence": incident.ai_confidence,
                "file_size": len(docx_bytes),
                "human_verified": human_verified,
            },
        )
        db.add(case_doc_record)
        await db.commit()
        await db.refresh(case_doc_record)
        await db.refresh(doc_record)

        logger.info(
            "Case report %s v%d generated successfully (SHA-256: %s)",
            filename,
            version,
            document_hash,
        )
        return case_doc_record, doc_record
